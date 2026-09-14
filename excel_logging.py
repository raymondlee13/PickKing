"""Spreadsheet logging -- appends scanned/selected legs to the tracking workbook."""

import os

from result_grading import build_tag, check_result, parse_tag

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# Best-effort stat abbreviations to match the workbook's existing convention
# (PTS, AST, etc). Falls back to the raw market key uppercased if unmapped.
STAT_ABBREV = {
    "player_points": "PTS", "player_rebounds": "REB", "player_assists": "AST",
    "player_points_rebounds_assists": "PRA", "player_threes": "3PM",
    "player_blocks": "BLK", "player_steals": "STL",
    "batter_hits": "HITS", "batter_home_runs": "HR", "batter_rbis": "RBI",
    "batter_hits_runs_rbis": "H+R+RBI", "batter_runs_scored": "RUNS",
    "batter_stolen_bases": "SB", "batter_total_bases": "TB",
    "pitcher_strikeouts": "K", "pitcher_hits_allowed": "H ALLOWED",
    "pitcher_walks": "BB", "pitcher_outs": "OUTS",
}
LINE_TYPE_MAP = {"standard": "STD", "goblin": "GOBLIN", "demon": "DEMON"}
TIER_CODE_MAP = {"TIER S": "S", "TIER A": "A", "TIER B": "B", "TIER C": "C", "BELOW BAR": "BELOW BAR"}


def log_parlay_to_excel(workbook_path, legs, multiplier, entry_type_label, date_str):
    """
    legs: list of dicts with player/market/point/side/dfs_type/consensus_pct/bar/margin/
          tier/whole_number/estimate_type/book_point/over_price/under_price/books/matchup/spread_pct
    Returns (success, message, entry_id).
    """
    if not OPENPYXL_AVAILABLE:
        return False, "openpyxl isn't installed. In Command Prompt run: pip install openpyxl --break-system-packages , then restart the app.", None
    if not workbook_path:
        return False, "No tracking workbook path set. Add \"tracking_workbook_path\" in config.json.", None
    if not os.path.exists(workbook_path):
        return False, f"Workbook not found at: {workbook_path}. Check the path in config.json.", None

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except PermissionError:
        return False, "Couldn't open the workbook -- close it in Excel first, then try again.", None
    except Exception as e:
        return False, f"Couldn't open workbook: {type(e).__name__}: {e}", None

    if "Leg Log" not in wb.sheetnames or "Entry Log" not in wb.sheetnames:
        return False, "Workbook doesn't have 'Leg Log' and 'Entry Log' sheets -- wrong file?", None

    leg_log = wb["Leg Log"]
    entry_log = wb["Entry Log"]

    # Next Entry ID: scan existing "E<number>" IDs across both sheets, increment past the max.
    existing_ids = set()
    for sheet in (leg_log, entry_log):
        for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            eid = row[1]
            if eid and isinstance(eid, str) and eid.startswith("E") and eid[1:].isdigit():
                existing_ids.add(int(eid[1:]))
    entry_id = f"E{(max(existing_ids) + 1) if existing_ids else 1}"

    combined_prob = 1.0
    # Find the first genuinely empty row rather than blindly appending after
    # max_row -- the sheet has hundreds of pre-formatted template rows with
    # formulas already in place but no actual data, and we want to fill those
    # in order (reusing their existing formulas) instead of skipping past all
    # of them and starting a fresh block far below.
    next_row = 2
    while leg_log.cell(row=next_row, column=5).value not in (None, ""):  # column E = Player
        next_row += 1
    legs_written = 0

    for leg in legs:
        stat = STAT_ABBREV.get(leg.get("market", ""), leg.get("market", "").upper())
        line_type = LINE_TYPE_MAP.get(leg.get("dfs_type", "standard"), "STD")
        true_pct = leg.get("consensus_pct", 0) / 100.0
        bar_dec = leg.get("bar", 0) / 100.0
        combined_prob *= true_pct
        tier_code = TIER_CODE_MAP.get(leg.get("tier", ""), leg.get("tier", ""))

        note = f"Logged via Edge Finder app. Books: {', '.join(leg.get('books', []))}."
        if leg.get("spread_pct") is not None:
            note += f" Spread: {leg['spread_pct']}pts."
        else:
            note += " Single book."
        # Embed sport/event/market so a later "check results" pass can
        # re-fetch this exact game's box score directly -- see result_grading.py.
        note += " " + build_tag(leg.get("sport", ""), leg.get("event_id", ""), leg.get("market", ""))

        r = next_row  # for formula references below, matching the workbook's own formula pattern
        values = [
            date_str, entry_id, entry_type_label, multiplier,
            leg.get("player", ""), leg.get("matchup", ""), stat, (leg.get("side") or "").upper(),
            leg.get("point"), f'=IF(I{r}="","",IF(I{r}=INT(I{r}),"Y","N"))', line_type,
            leg.get("book_point"), leg.get("over_price"), leg.get("under_price"),
            true_pct, leg.get("estimate_type", ""), bar_dec,
            f'=IF(OR(O{r}="",Q{r}=""),"",O{r}-Q{r})',
            tier_code,
            None, None,  # Actual Value / Result -- filled in later
            f'=IF(U{r}="W",1,IF(U{r}="L",0,""))',
            f'=IF(V{r}="","",(O{r}-V{r})^2)',
            None,  # CLV Favorable? -- filled in later
            note,
        ]
        for col, val in enumerate(values, start=1):
            leg_log.cell(row=next_row, column=col, value=val)
        leg_log.cell(row=next_row, column=4).number_format = '0.0\\x'
        leg_log.cell(row=next_row, column=15).number_format = '0.0%'
        leg_log.cell(row=next_row, column=17).number_format = '0.0%'
        leg_log.cell(row=next_row, column=18).number_format = '0.0%'
        leg_log.cell(row=next_row, column=23).number_format = '0.0000'
        next_row += 1
        legs_written += 1

    modelled_ev = round(combined_prob * multiplier, 3)
    next_entry_row = 2
    while entry_log.cell(row=next_entry_row, column=2).value not in (None, ""):  # column B = Entry ID
        next_entry_row += 1
    entry_row = next_entry_row
    entry_values = [date_str, entry_id, entry_type_label, multiplier, legs_written,
                     modelled_ev, None, None, None, None, "Logged via Edge Finder app."]
    for col, val in enumerate(entry_values, start=1):
        entry_log.cell(row=entry_row, column=col, value=val)
    entry_log.cell(row=entry_row, column=4).number_format = '0.0\\x'

    try:
        wb.save(workbook_path)
    except PermissionError:
        return False, "Couldn't save -- the workbook is open in Excel. Close it and try again.", None
    except Exception as e:
        return False, f"Couldn't save workbook: {type(e).__name__}: {e}", None

    return True, f"Logged {legs_written} leg(s) as entry {entry_id}.", entry_id


def log_legs_for_tracking(workbook_path, legs, date_str):
    """Log legs individually for tier-accuracy tracking -- no real entry, no
    multiplier, just a record of each leg's tier/probability at the time it
    was logged. Writes to 'Leg Log' only (never 'Entry Log', since there's no
    combined real entry here). 'Actual Value'/'Result' stay blank exactly like
    log_parlay_to_excel's rows -- fill in W/L by hand once the game finishes,
    and the sheet's existing hit/squared-error formulas do the rest, so you
    can pivot hit-rate by Tier to see whether the tiers actually predict
    anything. IDs use a "T<n>" namespace (not "E<n>") so these rows are
    visually distinguishable from real placed entries in the spreadsheet.
    Returns (success, message).
    """
    if not OPENPYXL_AVAILABLE:
        return False, "openpyxl isn't installed. In Command Prompt run: pip install openpyxl --break-system-packages , then restart the app."
    if not workbook_path:
        return False, "No tracking workbook path set. Add \"tracking_workbook_path\" in config.json."
    if not os.path.exists(workbook_path):
        return False, f"Workbook not found at: {workbook_path}. Check the path in config.json."

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except PermissionError:
        return False, "Couldn't open the workbook -- close it in Excel first, then try again."
    except Exception as e:
        return False, f"Couldn't open workbook: {type(e).__name__}: {e}"

    if "Leg Log" not in wb.sheetnames:
        return False, "Workbook doesn't have a 'Leg Log' sheet -- wrong file?"

    leg_log = wb["Leg Log"]

    existing_ids = set()
    for row in leg_log.iter_rows(min_row=2, max_col=2, values_only=True):
        eid = row[1]
        if eid and isinstance(eid, str) and eid.startswith("T") and eid[1:].isdigit():
            existing_ids.add(int(eid[1:]))
    next_id_num = (max(existing_ids) + 1) if existing_ids else 1

    next_row = 2
    while leg_log.cell(row=next_row, column=5).value not in (None, ""):  # column E = Player
        next_row += 1
    legs_written = 0

    for leg in legs:
        stat = STAT_ABBREV.get(leg.get("market", ""), leg.get("market", "").upper())
        line_type = LINE_TYPE_MAP.get(leg.get("dfs_type", "standard"), "STD")
        true_pct = leg.get("consensus_pct", 0) / 100.0
        bar_dec = leg.get("bar", 0) / 100.0
        tier_code = TIER_CODE_MAP.get(leg.get("tier", ""), leg.get("tier", ""))
        entry_id = f"T{next_id_num}"
        next_id_num += 1

        note = f"Tracking-only, not a real placed entry -- logged via Edge Finder app. Books: {', '.join(leg.get('books', []))}."
        if leg.get("spread_pct") is not None:
            note += f" Spread: {leg['spread_pct']}pts."
        else:
            note += " Single book."
        note += " " + build_tag(leg.get("sport", ""), leg.get("event_id", ""), leg.get("market", ""))

        r = next_row
        values = [
            date_str, entry_id, "TRACKING", None,
            leg.get("player", ""), leg.get("matchup", ""), stat, (leg.get("side") or "").upper(),
            leg.get("point"), f'=IF(I{r}="","",IF(I{r}=INT(I{r}),"Y","N"))', line_type,
            leg.get("book_point"), leg.get("over_price"), leg.get("under_price"),
            true_pct, leg.get("estimate_type", ""), bar_dec,
            f'=IF(OR(O{r}="",Q{r}=""),"",O{r}-Q{r})',
            tier_code,
            None, None,  # Actual Value / Result -- filled in later
            f'=IF(U{r}="W",1,IF(U{r}="L",0,""))',
            f'=IF(V{r}="","",(O{r}-V{r})^2)',
            None,  # CLV Favorable? -- filled in later
            note,
        ]
        for col, val in enumerate(values, start=1):
            leg_log.cell(row=next_row, column=col, value=val)
        leg_log.cell(row=next_row, column=15).number_format = '0.0%'
        leg_log.cell(row=next_row, column=17).number_format = '0.0%'
        leg_log.cell(row=next_row, column=18).number_format = '0.0%'
        leg_log.cell(row=next_row, column=23).number_format = '0.0000'
        next_row += 1
        legs_written += 1

    try:
        wb.save(workbook_path)
    except PermissionError:
        return False, "Couldn't save -- the workbook is open in Excel. Close it and try again."
    except Exception as e:
        return False, f"Couldn't save workbook: {type(e).__name__}: {e}"

    return True, f"Logged {legs_written} leg(s) for tracking (no multiplier -- fill in Result later)."


def check_and_fill_results(workbook_path, api_key):
    """Scan 'Leg Log' for rows with a blank Result but real leg data, re-fetch
    each one's real box score via the [pk:sport=...,event=...,market=...] tag
    embedded in its Note column at logging time (see build_tag in
    result_grading.py), and fill in Actual Value/Result where the game has
    finished. Rows logged before this tag existed, or where the game hasn't
    finished yet, are left alone -- counted separately so the caller can
    report what happened rather than silently doing nothing for them.
    Returns (success, message).
    """
    if not OPENPYXL_AVAILABLE:
        return False, "openpyxl isn't installed. In Command Prompt run: pip install openpyxl --break-system-packages , then restart the app."
    if not workbook_path:
        return False, "No tracking workbook path set. Add \"tracking_workbook_path\" in config.json."
    if not os.path.exists(workbook_path):
        return False, f"Workbook not found at: {workbook_path}. Check the path in config.json."

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except PermissionError:
        return False, "Couldn't open the workbook -- close it in Excel first, then try again."
    except Exception as e:
        return False, f"Couldn't open workbook: {type(e).__name__}: {e}"

    if "Leg Log" not in wb.sheetnames:
        return False, "Workbook doesn't have a 'Leg Log' sheet -- wrong file?"

    leg_log = wb["Leg Log"]
    stats_cache = {}  # (sport, event_id) -> box score, shared across every row this pass
    graded = pending = no_data = no_tag = 0
    changed = False

    row = 2
    while leg_log.cell(row=row, column=5).value not in (None, ""):  # column E = Player
        player = leg_log.cell(row=row, column=5).value
        if leg_log.cell(row=row, column=21).value not in (None, ""):  # column U = Result
            row += 1
            continue  # already graded, or filled in by hand -- leave it alone

        tag = parse_tag(leg_log.cell(row=row, column=25).value)  # column Y = Note
        if not tag:
            no_tag += 1
            row += 1
            continue

        side = "More" if str(leg_log.cell(row=row, column=8).value or "").upper() == "MORE" else "Less"
        point = leg_log.cell(row=row, column=9).value  # column I = Line

        try:
            status, result, actual_value = check_result(
                tag["sport"], tag["event_id"], tag["market"], player, point, side, api_key, stats_cache)
        except Exception:
            pending += 1  # API hiccup for this one -- try again next pass, don't fail the whole run
            row += 1
            continue

        if status != "final":
            pending += 1
        elif result is None:
            no_data += 1
        else:
            leg_log.cell(row=row, column=20, value=actual_value)  # column T = Actual Value
            leg_log.cell(row=row, column=21, value=result)  # column U = Result
            graded += 1
            changed = True
        row += 1

    if changed:
        try:
            wb.save(workbook_path)
        except PermissionError:
            return False, "Couldn't save -- the workbook is open in Excel. Close it and try again."
        except Exception as e:
            return False, f"Couldn't save workbook: {type(e).__name__}: {e}"

    message = (f"Graded {graded} pick(s). {pending} still pending (game not finished yet). "
               f"{no_data} finished but couldn't match the player/stat. "
               f"{no_tag} too old to auto-check (logged before this feature existed).")
    return True, message
