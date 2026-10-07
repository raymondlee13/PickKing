"""Checks for result grading, entry settlement, duplicate skipping, the Wilson
interval, and the workbook layout guard.
Run: python test_tracking.py"""

import os
import tempfile

from calibration_report import wilson_interval
from excel_logging import (
    COL, ENTRY_LOG_HEADERS, LEG_LOG_HEADERS, OPENPYXL_AVAILABLE, leg_prices_path, log_legs_for_tracking, log_parlay_to_excel,
    open_workbook, settle_entries,
)
from result_grading import build_tag, grade_pick, parse_tag

BOX = [
    {"player_name": "A'ja Wilson", "stat_type": "points", "stat_value": 22},
    {"player_name": "A'ja Wilson", "stat_type": "rebounds", "stat_value": 10},
    {"player_name": "A'ja Wilson", "stat_type": "assists", "stat_value": 3},
]


def test_grade_pick_sides_and_push():
    g = lambda point, side: grade_pick("player_points", "basketball_wnba", BOX, "A'ja Wilson", point, side)
    assert g(21.5, "More") == "W" and g(22.5, "More") == "L"
    assert g(22.5, "Less") == "W" and g(21.5, "Less") == "L"
    assert g(22, "More") == "PUSH" and g(22, "Less") == "PUSH"  # whole-number line


def test_grade_pick_combined_and_missing():
    # PRA isn't a single stat_type -- it's the sum of three.
    assert grade_pick("player_points_rebounds_assists", "basketball_wnba", BOX, "A'ja Wilson", 34.5, "More") == "W"
    assert grade_pick("player_points_rebounds_assists", "basketball_wnba", BOX, "A'ja Wilson", 35, "More") == "PUSH"
    assert grade_pick("player_points", "basketball_wnba", BOX, "Someone Else", 10.5, "More") is None
    assert grade_pick("player_blocks", "basketball_wnba", BOX, "A'ja Wilson", 0.5, "More") is None  # stat absent
    assert grade_pick("not_a_market", "basketball_wnba", BOX, "A'ja Wilson", 0.5, "More") is None
    assert grade_pick("player_points", "basketball_wnba", BOX, "A'ja Wilson", None, "More") is None


def test_tag_round_trip():
    note = "Logged via app. " + build_tag("baseball_mlb", "abc123", "batter_hits")
    assert parse_tag(note) == {"sport": "baseball_mlb", "event_id": "abc123", "market": "batter_hits"}
    assert parse_tag("hand-typed note") is None and parse_tag(None) is None


def test_wilson_interval():
    assert wilson_interval(0, 0) is None
    low, high = wilson_interval(7, 11)
    assert 0.34 < low < 0.36 and 0.84 < high < 0.86
    low, high = wilson_interval(0, 5)  # textbook interval collapses to 0-0 here
    assert low == 0 and high > 0.4
    # More picks at the same rate -> narrower range.
    assert (lambda a, b: b[1] - b[0] < a[1] - a[0])(wilson_interval(7, 11), wilson_interval(64, 100))


def test_open_workbook_rejects_moved_columns():
    if not OPENPYXL_AVAILABLE:
        return
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Leg Log"
        wb.active.append(LEG_LOG_HEADERS)
        wb.save(path)
        assert open_workbook(path)[1] is None

        wb.active.insert_cols(3)  # someone adds a column in Excel
        wb.save(path)
        _, error = open_workbook(path)
        assert error and "column 3" in error
        assert open_workbook(os.path.join(d, "missing.xlsx"))[1].startswith("Workbook not found")


def _blank_workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Leg Log"
    wb.active.append(LEG_LOG_HEADERS)
    wb.create_sheet("Entry Log").append(ENTRY_LOG_HEADERS)
    wb.save(path)


def test_settle_entries():
    if not OPENPYXL_AVAILABLE:
        return
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        _blank_workbook(path)
        wb = openpyxl.load_workbook(path)
        leg_log, entry_log = wb["Leg Log"], wb["Entry Log"]
        # entry id, type, multiplier, stake, notes, leg results
        entries = [
            ("E1", "2-Power", 3.0, 10, "", ["W", "W"]),
            ("E2", "3-Power", 5.0, None, "", ["W", "L", "W"]),
            ("E3", "4-Flex", 6.0, 10, "[pkflex:1miss=1.5,2miss=]", ["W", "L", "W", "W"]),
            ("E4", "5-Flex", 10.0, None, "[pkflex:1miss=2,2miss=0.4]", ["L", "L", "L", "W", "W"]),
            ("E5", "2-Power", 3.0, None, "", ["W", "PUSH"]),
            ("E6", "2-Power", 3.0, None, "", ["W", None]),
        ]
        for eid, typ, mult, stake, notes, results in entries:
            entry_log.append(["2026-10-01", eid, typ, mult, len(results), 1.1, stake, None, None, None, notes])
            for res in results:
                row = [None] * len(LEG_LOG_HEADERS)
                row[COL["Entry ID"] - 1], row[COL["Player"] - 1], row[COL["Result"] - 1] = eid, "P", res
                leg_log.append(row)

        assert settle_entries(wb) == (4, 1)  # E5 pushed -> manual, E6 still pending
        got = {entry_log.cell(row=r, column=2).value: (entry_log.cell(row=r, column=8).value,
                                                        entry_log.cell(row=r, column=9).value)
               for r in range(2, 8)}
        assert got["E1"] == ("W", 30.0)          # 10 staked x 3.0
        assert got["E2"] == ("L", None)          # Power: one miss loses; no stake -> no Return
        assert got["E3"] == ("Partial", 15.0)    # Flex 1 miss pays its 1.5x tier
        assert got["E4"] == ("L", None)          # 3 misses: below every Flex tier
        assert got["E5"] == (None, None) and got["E6"] == (None, None)
        assert settle_entries(wb) == (0, 1)      # already-settled rows are left alone


def test_tracking_log_skips_duplicates():
    if not OPENPYXL_AVAILABLE:
        return
    import openpyxl
    leg = {"player": "A'ja Wilson", "market": "player_points", "point": 21.5, "side": "More",
           "dfs_type": "standard", "consensus_pct": 60.0, "bar": 55.0, "tier": "TIER B",
           "books": ["fanduel"], "sport": "basketball_wnba", "event_id": "123"}
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        _blank_workbook(path)
        assert log_legs_for_tracking(path, [leg, dict(leg)], "2026-10-01")[0]  # dup within one batch
        ok, message = log_legs_for_tracking(path, [leg, dict(leg, point=22.5)], "2026-10-01")
        assert ok and "Skipped 1" in message
        leg_log = openpyxl.load_workbook(path)["Leg Log"]
        assert [leg_log.cell(row=r, column=COL["PP Line"]).value for r in (2, 3, 4)] == [21.5, 22.5, None]


def test_logged_stake_settles_to_return():
    if not OPENPYXL_AVAILABLE:
        return
    import openpyxl
    leg = {"player": "A'ja Wilson", "market": "player_points", "point": 21.5, "side": "More",
           "dfs_type": "standard", "consensus_pct": 60.0, "bar": 55.0, "tier": "TIER B",
           "books": ["fanduel"], "sport": "basketball_wnba", "event_id": "123"}
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        _blank_workbook(path)
        assert log_parlay_to_excel(path, [leg, dict(leg, point=4.5)], 3.0, "2-Power", "2026-10-01", stake=5)[0]
        wb = openpyxl.load_workbook(path)
        for r in (2, 3):  # what "Check results" writes once both games are final
            wb["Leg Log"].cell(row=r, column=COL["Result"], value="W")
        assert settle_entries(wb) == (1, 0)
        entry_log = wb["Entry Log"]
        assert [entry_log.cell(row=2, column=c).value for c in (7, 8, 9)] == [5, "W", 15.0]  # Stake, Result, Return


def test_logging_saves_every_books_price():
    if not OPENPYXL_AVAILABLE:
        return
    import json
    leg = {"player": "A'ja Wilson", "market": "player_points", "point": 21.5, "side": "More",
           "dfs_type": "standard", "consensus_pct": 60.0, "bar": 55.0, "tier": "TIER B", "books": ["fanduel"],
           "sport": "basketball_wnba", "event_id": "123",
           "book_prices": [["fanduel", 21.5, -150, 120], ["novig", 21.5, -140, 130]]}
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        _blank_workbook(path)
        log_legs_for_tracking(path, [leg], "2026-10-01")
        log_legs_for_tracking(path, [leg], "2026-10-01")  # duplicate: skipped, so no second line
        log_parlay_to_excel(path, [leg, dict(leg, point=4.5)], 3.0, "2-Power", "2026-10-01")
        with open(leg_prices_path(path), encoding="utf-8") as f:
            lines = [json.loads(x) for x in f]
        assert [x["entry_id"] for x in lines] == ["T1", "E1", "E1"]
        assert lines[0]["book_prices"] == leg["book_prices"] and lines[2]["point"] == 4.5


def test_correlation_baseline_uses_each_pairs_probabilities():
    if not OPENPYXL_AVAILABLE:
        return
    import openpyxl
    from correlation_report import build_report
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        _blank_workbook(path)
        wb = openpyxl.load_workbook(path)
        for player, true_pct, result in (("A", 0.9, "W"), ("B", 0.5, "W"), ("C", 0.5, "L")):
            row = [None] * len(LEG_LOG_HEADERS)
            for h, v in (("Date", "2026-10-01"), ("Game", "X @ Y"), ("Player", player), ("PP Line", 1.5),
                         ("Side", "MORE"), ("Predicted True %", true_pct), ("Result", result)):
                row[COL[h] - 1] = v
            wb["Leg Log"].append(row)
        wb.save(path)
        summary = build_report(path)[2]["summary"]
        # Pairs: A-B 0.45, A-C 0.45, B-C 0.25 -> 38.3%. Old baseline: (2/3)^2 = 44.4%.
        assert summary["expected_both_hit_pct_if_independent"] == 38.3


def test_daily_backup_once_and_pruned():
    import datetime
    import excel_logging
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "Track.xlsx")
        with open(path, "w") as f:
            f.write("v1")
        folder = os.path.join(d, "backups")
        os.makedirs(folder)
        for day in range(1, 21):  # 20 old backups already there
            open(os.path.join(folder, f"Track-2026-01-{day:02d}.xlsx"), "w").close()
        excel_logging._backup_workbook_today(path)
        with open(path, "w") as f:
            f.write("v2")
        excel_logging._backup_workbook_today(path)  # second write today: no new copy
        today = os.path.join(folder, f"Track-{datetime.date.today().isoformat()}.xlsx")
        with open(today) as f:
            assert f.read() == "v1"  # holds the state from before the day's first write
        assert len(os.listdir(folder)) == excel_logging.BACKUPS_KEPT and os.path.exists(today)


def test_clv_check_records_closing_prices():
    if not OPENPYXL_AVAILABLE:
        return
    import json
    import openpyxl
    import excel_logging
    closing = {"bookmakers": [{"key": "fanduel", "markets": [{"key": "player_points", "outcomes": [
        {"name": "Over", "description": "A'ja Wilson", "point": 21.5, "price": -170},
        {"name": "Under", "description": "A'ja Wilson", "point": 21.5, "price": 140}]}]}]}

    def fake_check_clv(sport, event_id, market, player, point, side, api_key, cache):
        cache[(sport, event_id, "odds", market)] = closing
        return "live", True, 22.5

    saved = excel_logging.check_clv
    excel_logging.check_clv = fake_check_clv
    try:
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "t.xlsx")
            _blank_workbook(path)
            leg = {"player": "A'ja Wilson", "market": "player_points", "point": 21.5, "side": "More",
                   "dfs_type": "standard", "consensus_pct": 60.0, "bar": 55.0, "tier": "TIER B",
                   "books": ["fanduel"], "sport": "basketball_wnba", "event_id": "123"}
            log_legs_for_tracking(path, [leg], "2026-10-01")
            assert excel_logging.check_and_fill_clv(path, "key")[0]
            assert openpyxl.load_workbook(path)["Leg Log"].cell(row=2, column=COL["CLV Favorable?"]).value == "YES"
            with open(leg_prices_path(path), encoding="utf-8") as f:
                records = [json.loads(x) for x in f]
            assert [r["kind"] for r in records] == ["log", "close"]
            assert records[1]["book_prices"] == [["fanduel", 21.5, -170, 140]] and records[1]["closing_point"] == 22.5
            assert records[1]["entry_id"] == records[0]["entry_id"] == "T1"
    finally:
        excel_logging.check_clv = saved


def test_deleted_user_loses_session():
    import auth
    with tempfile.TemporaryDirectory() as d:
        saved = auth.USERS_PATH, auth.SESSIONS_PATH, dict(auth._sessions)
        auth.USERS_PATH, auth.SESSIONS_PATH = os.path.join(d, "users.json"), os.path.join(d, "sessions.json")
        try:
            assert auth.create_user("ray", "password123")[0]
            token = auth.create_session("ray")
            assert auth.username_for_session(token) == "ray"
            os.remove(auth.USERS_PATH)  # account removed by hand
            assert auth.username_for_session(token) is None
        finally:
            auth.USERS_PATH, auth.SESSIONS_PATH = saved[0], saved[1]
            auth._sessions.clear()
            auth._sessions.update(saved[2])


if __name__ == "__main__":
    tests = [f for name, f in sorted(globals().items()) if name.startswith("test_")]
    for t in tests:
        t()
    print(f"{len(tests)} checks passed")
