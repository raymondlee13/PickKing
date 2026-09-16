"""Reads the tracking workbook's 'Leg Log' sheet and reports on same-game
leg pairs -- do legs from the same real game tend to hit or miss together
more than you'd expect if their outcomes were independent?

This answers a genuinely different question than the "Check correlation"
AI feature (see ai_context.py). That gives a one-time qualitative opinion
*before* you bet, based on the model's general sports knowledge, and it's
never checked against what actually happened. This module instead looks
backward at REAL logged results: for every pair of legs you've logged that
came from the same game on the same date, what actually happened once both
resolved? That's empirical evidence, not a guess -- but it needs real
sample size to mean anything. A handful of pairs here should be read as
"not enough data yet," not "no correlation" -- this is designed to
accumulate value over weeks of use, not on day one.

Groups by (date, matchup) rather than requiring the legs to have been part
of the same logged entry -- a discount/tracking-only leg logged separately
from a real parlay leg for the same game still counts, since what matters
for this question is just "were these two legs riding on the same game,"
not how or when they were logged.
"""

import os

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _read_graded_rows(leg_log):
    rows = []
    r = 2
    while leg_log.cell(row=r, column=5).value not in (None, ""):  # column E = Player
        result = leg_log.cell(row=r, column=21).value  # column U = Result
        if result in ("W", "L"):  # exclude blank/pending and PUSH -- neither is a real outcome
            rows.append({
                "date": leg_log.cell(row=r, column=1).value,
                "matchup": leg_log.cell(row=r, column=6).value,
                "player": leg_log.cell(row=r, column=5).value,
                "stat": leg_log.cell(row=r, column=7).value,
                "point": leg_log.cell(row=r, column=9).value,
                "side": leg_log.cell(row=r, column=8).value,
                "hit": result == "W",
            })
        r += 1
    return rows


def build_report(workbook_path):
    """Returns (success, message, data). data is None on failure, otherwise:
    {total_legs, overall_hit_rate, pairs, summary}
    """
    if not OPENPYXL_AVAILABLE:
        return False, "openpyxl isn't installed. In Command Prompt run: pip install openpyxl --break-system-packages , then restart the app.", None
    if not workbook_path:
        return False, "No tracking workbook path set. Add \"tracking_workbook_path\" in config.json.", None
    if not os.path.exists(workbook_path):
        return False, f"Workbook not found at: {workbook_path}. Check the path in config.json.", None

    try:
        wb = openpyxl.load_workbook(workbook_path)
    except PermissionError:
        return False, "Couldn't open the workbook -- close it in Excel first, then try again.", None
    except Exception as e:
        return False, f"Couldn't open workbook: {type(e).__name__}: {e}", None

    if "Leg Log" not in wb.sheetnames:
        return False, "Workbook doesn't have a 'Leg Log' sheet -- wrong file?", None

    rows = _read_graded_rows(wb["Leg Log"])
    if not rows:
        return True, "No graded picks yet (Result = W/L) -- log some picks and run \"Check results\" first.", {
            "total_legs": 0, "overall_hit_rate": None, "pairs": [], "summary": None,
        }

    overall_hit_rate = sum(1 for r in rows if r["hit"]) / len(rows)

    groups = {}
    for r in rows:
        key = (r["date"], r["matchup"])
        if not r["matchup"]:
            continue  # no matchup recorded (e.g. very old rows) -- can't group these
        groups.setdefault(key, []).append(r)

    pairs = []
    for (date, matchup), legs in groups.items():
        for i in range(len(legs)):
            for j in range(i + 1, len(legs)):
                a, b = legs[i], legs[j]
                if a["hit"] and b["hit"]:
                    outcome = "BOTH_HIT"
                elif not a["hit"] and not b["hit"]:
                    outcome = "BOTH_MISS"
                else:
                    outcome = "SPLIT"
                pairs.append({
                    "date": date, "matchup": matchup,
                    "leg_a": f"{a['player']} {a['side']} {a['stat']} {a['point']}",
                    "leg_b": f"{b['player']} {b['side']} {b['stat']} {b['point']}",
                    "outcome": outcome,
                })

    summary = None
    if pairs:
        total = len(pairs)
        both_hit = sum(1 for p in pairs if p["outcome"] == "BOTH_HIT")
        both_miss = sum(1 for p in pairs if p["outcome"] == "BOTH_MISS")
        split = total - both_hit - both_miss
        # Rough independence baseline: if two legs' outcomes were unrelated,
        # both-hit rate would be roughly (your overall hit rate)^2. Approximate
        # on purpose -- a real baseline would use each leg's own probability,
        # not one global rate, but that needs more data than this tool will
        # realistically have early on. Treat this as a sanity-check reference
        # point, not a rigorous statistical test.
        summary = {
            "total_pairs": total,
            "both_hit_pct": round(both_hit / total * 100, 1),
            "both_miss_pct": round(both_miss / total * 100, 1),
            "split_pct": round(split / total * 100, 1),
            "expected_both_hit_pct_if_independent": round((overall_hit_rate ** 2) * 100, 1),
        }

    data = {
        "total_legs": len(rows),
        "overall_hit_rate": round(overall_hit_rate * 100, 1),
        "pairs": pairs,
        "summary": summary,
    }
    return True, "OK", data
